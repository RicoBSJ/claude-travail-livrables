#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""refaire_classeur.py — vérifie, ou refait, le classeur de suivi des jobs.

POURQUOI (15/10/2026). Le classeur `livrables/documents/*_Les N Jobs.xlsx` est le SEUL
endroit où un humain lit la liste des jobs, et c'était le seul que rien ne vérifiait :
`controle_clauses_parcours.py` compare le prompt aux `.docx`, `audit_vault.py` les fiches
aux documents, et **personne ne comparait le classeur à `jobs_config.json`**.

Résultat mesuré ce jour-là : `2026-08-09_Les 14 Jobs.xlsx` portait encore CINQ jobs
supprimés et s'intitulait « Tes 14 jobs » pour 9 jobs réels. La cause n'est pas un oubli
isolé : `maj_xlsx_jobs.py` écrit dans le classeur après CHAQUE job réussi, mais il ne met à
jour QU'UNE COLONNE — « Numéro de la dernière leçon » — et ne touche jamais au jeu de
lignes. **Il est donc structurellement incapable de suivre une suppression de job**, et le
classeur dérive en silence à chaque fois. Refait à la main le 15/10 en « Les 9 Jobs » ; ce
script généralise ce geste pour qu'il ne se refasse plus à la main.

Deux modes :
  --verifier   compare le classeur à jobs_config.json. exit 0 s'ils concordent, 1 sinon,
               en nommant l'écart. Aucun réseau, instantané : c'est le contrôle manquant.
  --refaire    construit `<aujourd'hui>_Les N Jobs.xlsx` depuis jobs_config.json, archive
               l'ancien dans documents/archives/, et REPOINTE `CLASSEUR` dans
               maj_xlsx_jobs.py — sans quoi celui-ci continuerait d'écrire dans le fichier
               archivé et le nouveau classeur se figerait le jour de sa création.

⚠️ Le classeur est construit DEPUIS la config, jamais recopié : les créneaux sont déduits
des expressions `cron`. La colonne « Type » est REPORTÉE de l'ancien classeur pour ne rien
inventer, et le script REFUSE si un job n'y a pas de type — plutôt que d'en fabriquer un.
Les numéros de leçon sont ensuite remplis par `maj_xlsx_jobs.py`, qui les LIT sur le disque.
"""
import datetime, json, re, shutil, subprocess, sys
from copy import copy
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent.parent
CONFIG = RACINE / "jobs_config.json"
DOCS = RACINE / "livrables" / "documents"
MAJ = RACINE / "outils" / "scripts" / "maj_xlsx_jobs.py"
JOURS = ["Dimanche", "Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi"]


def classeur_actuel():
    """Le classeur vivant : celui que maj_xlsx_jobs.py pointe, qui est la seule référence."""
    m = re.search(r'CLASSEUR = RACINE / "livrables" / "documents" / "([^"]+)"',
                  MAJ.read_text(encoding="utf-8"))
    if not m:
        sys.exit("⛔ CLASSEUR introuvable dans maj_xlsx_jobs.py")
    return DOCS / m.group(1)


def jobs_config():
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))["jobs"]
    out = []
    for j in cfg:
        mi, h, _, _, wd = j["cron"].split()
        out.append({"id": j["id"], "wd": int(wd), "h": int(h), "mi": int(mi),
                    "jour": JOURS[int(wd)], "heure": "%dh%02d" % (int(h), int(mi))})
    return sorted(out, key=lambda x: (x["wd"], x["h"], x["mi"]))


def lignes_classeur(f):
    import openpyxl
    fe = openpyxl.load_workbook(f).active
    ent = None
    for lg in fe.iter_rows(min_row=1, max_row=6):
        for c in lg:
            if c.value == "Job":
                ent, col = c.row, c.column
                break
        if ent:
            break
    if not ent:
        sys.exit("⛔ colonne « Job » introuvable dans %s" % f.name)
    ids, types = [], {}
    for r in range(ent + 1, fe.max_row + 1):
        v = fe.cell(row=r, column=col).value
        if v:
            ids.append(str(v).strip())
            types[str(v).strip()] = fe.cell(row=r, column=col + 1).value
    return fe, ent, col, ids, types


def verifier():
    f = classeur_actuel()
    if not f.exists():
        print("⛔ CLASSEUR pointe un fichier absent : %s" % f.name)
        return 1
    reels = [j["id"] for j in jobs_config()]
    _, _, _, ids, _ = lignes_classeur(f)
    titre = f.name
    pb = []
    manquants = [i for i in reels if i not in ids]
    morts = [i for i in ids if i not in reels]
    if morts:
        pb.append("%d job(s) SUPPRIMÉ(S) encore listé(s) : %s" % (len(morts), ", ".join(morts)))
    if manquants:
        pb.append("%d job(s) actif(s) ABSENT(S) : %s" % (len(manquants), ", ".join(manquants)))
    m = re.search(r"_Les (\d+) Jobs", titre)
    if m and int(m.group(1)) != len(reels):
        pb.append("le titre du fichier annonce %s jobs pour %d réels" % (m.group(1), len(reels)))
    print("▶ Classeur : %s" % titre)
    print("  %d ligne(s) de job · %d job(s) dans jobs_config.json" % (len(ids), len(reels)))
    if pb:
        for x in pb:
            print("  ✗ %s" % x)
        print("  → `refaire_classeur.py --refaire` reconstruit le classeur depuis la config.")
        return 1
    print("  ✓ le classeur concorde avec jobs_config.json")
    return 0


def refaire():
    import openpyxl
    ancien = classeur_actuel()
    reels = jobs_config()
    _, ent, col, _, types = lignes_classeur(ancien)
    sans = [j["id"] for j in reels if not types.get(j["id"])]
    if sans:
        sys.exit("⛔ aucun « Type » dans l'ancien classeur pour : %s\n"
                 "   Je n'en invente pas : ajoute-les à la main, puis relance." % ", ".join(sans))
    fv = openpyxl.load_workbook(ancien).active
    neuf = openpyxl.Workbook()
    f = neuf.active
    f.title = fv.title
    f.cell(1, 1, "Tes %d jobs" % len(reels))._style = copy(fv.cell(1, 1)._style)
    for c in range(1, fv.max_column + 1):
        f.cell(2, c, fv.cell(ent, c).value)._style = copy(fv.cell(ent, c)._style)
    jour_vu = None
    for i, j in enumerate(reels):
        r = 3 + i
        vals = [j["jour"] if j["jour"] != jour_vu else None, j["heure"], j["id"],
                types[j["id"]], "—"]
        jour_vu = j["jour"]
        modele = min(ent + 1 + i, fv.max_row)
        for c, v in enumerate(vals[:fv.max_column], 1):
            f.cell(r, c, v)._style = copy(fv.cell(modele, c)._style)
    for k, d in fv.column_dimensions.items():
        if d.width:
            f.column_dimensions[k].width = d.width
    cible = DOCS / ("%s_Les %d Jobs.xlsx" % (datetime.date.today().isoformat(), len(reels)))
    if cible.resolve() == ancien.resolve():
        sys.exit("⛔ le nom cible est celui du classeur actuel : rien à refaire aujourd'hui")
    neuf.save(cible)
    print("  ✓ écrit %s (%d jobs)" % (cible.name, len(reels)))
    (DOCS / "archives").mkdir(exist_ok=True)
    shutil.move(str(ancien), str(DOCS / "archives" / ancien.name))
    print("  ✓ archivé %s → documents/archives/" % ancien.name)
    s = MAJ.read_text(encoding="utf-8")
    MAJ.write_text(s.replace('"%s"' % ancien.name, '"%s"' % cible.name, 1), encoding="utf-8")
    print("  ✓ CLASSEUR repointé dans maj_xlsx_jobs.py")
    r = subprocess.run([sys.executable, str(MAJ), "--verbeux"], capture_output=True, text=True, env={})
    print("  ✓ numéros de leçon remplis depuis le disque" if r.returncode == 0
          else "  ⚠ maj_xlsx_jobs.py a rendu %d : %s" % (r.returncode, r.stdout[-200:]))
    print("  ⚠ git mv n'a pas été utilisé : pense à `git add` l'ancien ET le nouveau chemin.")
    return 0


if __name__ == "__main__":
    a = sys.argv[1:]
    if a == ["--verifier"]:
        sys.exit(verifier())
    if a == ["--refaire"]:
        sys.exit(refaire())
    sys.exit(__doc__)
